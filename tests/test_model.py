import unittest
from dataclasses import replace
import numpy as np
import pandas as pd
from endowment_simulation_model import *
from portfolio_return_simulation import compare_return_methods
from strategy_analysis import portfolio_from_assets, policy_cases, compare_cases
from stress_testing import stress_cases, mixture_sensitivity
from validation import holdout_diagnostics, stability_check

class ModelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data=load_default_data()
        cls.c=SimulationConfig(initial_endowment=100,prior_endowment_1y=100,prior_endowment_2y=100,horizon_years=3,n_sims=12,return_method='Single Normal')
        cls.i=InstitutionalParams(effective_spending_rate=.05,gift_rate=0,other_changes_rate=0)
        cls.l=LiquidityParams()
    def path(self,returns,liq=None,inst=None,config=None):
        c=config or replace(self.c,horizon_years=len(returns))
        market=pd.DataFrame({'year':np.arange(len(returns))+2027,'portfolio_return':returns,'inflation':0.})
        return run_path(market,c,inst or self.i,liq or self.l)
    def test_asset_depletion_does_not_fund_impossible_payout(self):
        p=self.path([-.99,0,0]);self.assertAlmostEqual(p.actual_payout.iloc[0],1)
        self.assertTrue((p.actual_payout.iloc[1:]==0).all())
        self.assertTrue((p.payout_shortfall_indicator==1).all())
    def test_rollforward_and_bucket_totals_under_forced_sales(self):
        p=self.path([-.1,-.2,.1],replace(self.l,enabled=True,liquid_share=.02,illiquid_share=.8,
            forced_sale_discount=.3,other_annual_cash_need=2,moderately_liquid_access=.5))
        self.assertTrue((p.accounting_residual.abs()<1e-10).all())
        np.testing.assert_allclose(p.ending_assets,p.ending_liquid_assets+p.ending_moderately_liquid_assets+p.ending_illiquid_assets)
        self.assertTrue((p.forced_sale_loss<=p.forced_sale).all())
    def test_no_liquidity_bucket_reset(self):
        p=self.path([0,0],replace(self.l,enabled=True,liquid_share=.01,illiquid_share=.99,forced_sale_discount=.25))
        self.assertAlmostEqual(p.illiquid_assets.iloc[1],p.ending_illiquid_assets.iloc[0])
        self.assertLess(p.ending_illiquid_assets.iloc[1],p.ending_illiquid_assets.iloc[0])
    def test_unpaid_other_obligation_is_explicit(self):
        p=self.path([0],replace(self.l,enabled=True,liquid_share=.01,illiquid_share=.99,allow_forced_sales=False,other_annual_cash_need=10))
        self.assertAlmostEqual(p.other_cash_paid.iloc[0],1)
        self.assertAlmostEqual(p.unpaid_other_cash_need.iloc[0],9)
        self.assertEqual(p.actual_payout.iloc[0],0)
    def test_moderately_liquid_bucket_access(self):
        p=self.path([0],replace(self.l,enabled=True,liquid_share=0,illiquid_share=0,moderately_liquid_access=1))
        self.assertAlmostEqual(p.actual_payout.iloc[0],5)
        self.assertAlmostEqual(p.ending_moderately_liquid_assets.iloc[0],95)
    def test_return_coverage_shortfall_can_have_full_payout(self):
        p=self.path([0]);self.assertEqual(p.return_coverage_shortfall_indicator.iloc[0],1)
        self.assertEqual(p.payout_shortfall_indicator.iloc[0],0)
        self.assertAlmostEqual(p.return_coverage_shortfall_amount.iloc[0],5)
    def test_reserve_priority(self):
        p=self.path([0],replace(self.l,enabled=True,liquid_share=.06,illiquid_share=.94,
            allow_forced_sales=False,liquidity_buffer_years=1))
        self.assertAlmostEqual(p.actual_payout.iloc[0],1)
        self.assertAlmostEqual(p.ending_liquid_assets.iloc[0],5)
    def test_spells_and_censored_recovery(self):
        self.assertEqual(longest_spell([1,1,0,1,1,1]),3)
        p=pd.DataFrame({'real_corpus_below_threshold':[0,1,1,0]})
        self.assertEqual(recovery_statistics(p)['recovery_years'],2)
        p=pd.DataFrame({'real_corpus_below_threshold':[0,1,1]})
        self.assertTrue(recovery_statistics(p)['recovery_censored'])
        self.assertTrue(np.isnan(recovery_statistics(p)['recovery_years']))
    def test_current_assets_policy(self):
        p=self.path([.1,0],inst=replace(self.i,spending_rule='Current assets'))
        self.assertAlmostEqual(p.planned_payout.iloc[1],5.25)
    def test_collar(self):
        p=self.path([.5,0],inst=replace(self.i,annual_payout_growth_cap=0))
        self.assertEqual(p.planned_payout.iloc[1],5)
    def test_annual_compounding(self):
        q=pd.DataFrame({'p4_return':[.1]*4,'cpi_yoy':[.02]*4,'condition':['Normal']*4})
        self.assertAlmostEqual(quarterly_to_annual(q,2027).portfolio_return.iloc[0],1.1**4-1)
    def test_invalid_inputs_rejected(self):
        for l in [replace(self.l,liquid_share=.9,illiquid_share=.9),replace(self.l,forced_sale_discount=1)]:
            with self.assertRaises(ValueError):self.path([0],l)
        with self.assertRaises(ValueError): self.path([0],config=replace(self.c,horizon_years=1,adverse_probability=1.1))
    def test_tail_es(self):
        from endowment_simulation_model import _tail_var_es
        var,es=_tail_var_es(np.arange(100)/100,.95)
        self.assertAlmostEqual(es,np.mean(np.arange(95,100)/100))
        self.assertGreaterEqual(es,var)
    def test_all_return_generators_reproducible_and_metrics_without_paths(self):
        for method in ['Single Normal','Two-component Normal mixture','Quarterly block bootstrap','Regime-switching bootstrap']:
            c=replace(self.c,return_method=method)
            a=run_monte_carlo(self.data,c,self.i,self.l)
            b=run_monte_carlo(self.data,c,self.i,self.l,keep_paths=False)
            pd.testing.assert_frame_equal(a['path_summary'],b['path_summary'])
            self.assertEqual(a['metrics'],b['metrics'])
    def test_integration_comparisons(self):
        c=replace(self.c,n_sims=3,horizon_years=2)
        table,_,_,_=compare_return_methods(self.data,c,self.i,self.l)
        self.assertEqual(len(table.columns),3)
        self.assertEqual(len(compare_cases(self.data,policy_cases(c,self.i,self.l))),4)
        self.assertEqual(len(compare_cases(self.data,stress_cases(c,self.i,self.l))),6)
        self.assertEqual(len(mixture_sensitivity(self.data,c,self.i,self.l)),6)
        self.assertEqual(len(stability_check(self.data,c,self.i,self.l,seeds=(1,2),counts=(2,3))),4)
        self.assertEqual(len(holdout_diagnostics(self.data.p4_quarterly)),2)
    def test_allocation_calibration_changes_with_weights(self):
        q=self.data.p4_quarterly.quarter.iloc[:12].to_numpy()
        assets=pd.DataFrame({'quarter':q,'A':np.linspace(-.1,.2,12),'B':np.linspace(-.02,.01,12)})
        wa=pd.DataFrame({'asset':['A','B'],'weight':[.8,.2],'liquidity_class':['High','Low']})
        wb=wa.copy();wb['weight']=[.2,.8]
        ha,la,ia=portfolio_from_assets(assets,wa,self.data.condition_history)
        hb,lb,ib=portfolio_from_assets(assets,wb,self.data.condition_history)
        self.assertNotEqual(ha.p4_return.std(),hb.p4_return.std())
        self.assertAlmostEqual(la,.8);self.assertAlmostEqual(ib,.8)
    def test_zero_excel_overrides_preserved(self):
        from openpyxl import Workbook
        from io import BytesIO
        w=Workbook();w.active.title='Model Inputs';w.create_sheet('R2 Inputs');w.create_sheet('R1 Inputs')
        for cell,val in {'B8':100,'B9':100,'B10':100,'B22':0.,'B24':1.,'B31':0,'B32':0,'B33':0}.items():w['Model Inputs'][cell]=val
        for j in range(5,11):w['R2 Inputs'][f'B{j}']=0
        f=BytesIO();w.save(f)
        _,i,l,_=load_excel_overrides(f.getvalue(),self.data)
        self.assertEqual(i.effective_spending_rate,0)
        self.assertEqual(l.liquid_share,0)

if __name__=='__main__': unittest.main()
